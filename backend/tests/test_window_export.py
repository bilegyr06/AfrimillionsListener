"""Campaign Window Excel evaluation export tests (v2 segment export).

Covers app.services.window_export and app.services.window_report.segment_evaluations
plus the GET /windows/{id}/export endpoint:

  * workbook structure (blocks, evaluation title, window line, metric headers,
    Campaign/Control rows, blank separator rows);
  * segment separation (one / multiple segments, multiple Runs, repeated same
    segment, no double counting, no merging across segments);
  * Campaign/Control values from the persisted assignment; the export is read-only;
  * representative values for every exported metric;
  * edge cases: empty groups, zero activity, `None` rates rendered as an em dash
    (never zero), multiple Runs, active/ended/finalized Windows, empty audience.

Timing and deterministic-assignment conventions mirror test_window_report.py:
window Mon 2026-09-14 -> Sat 2026-09-19 23:59:59, Sun 2026-09-20 14:00 deadline;
control override 50% -> campaign_percentage 50, so user_id % 100 < 50 is Campaign.
"""
from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from app.core import dates
from app.db.deposits import insert_deposit_records
from app.db.logins import insert_login_records
from app.db.players import insert_play_records
from app.db.sms import log_sms
from app.services import window_export, window_report
from app.services import windows as svc
from app.services.logins import login_source_key

LAGOS = ZoneInfo("Africa/Lagos")

GROUP_HEADER = window_export.GROUP_HEADER
METRIC_HEADERS = [header for header, _key, _fmt in window_export.METRIC_COLUMNS]
UNAVAILABLE = window_export.UNAVAILABLE

CAMPAIGN_IDS = ("1", "2", "3")
CONTROL_IDS = ("51", "52")

FINALIZE_AT = datetime(2026, 9, 20, 9, 0, 0, tzinfo=LAGOS)


@pytest.fixture()
def client(_init_db):
    from fastapi.testclient import TestClient
    from app.main import app
    return TestClient(app)


def _dt(y, mo, d, h=0, mi=0, s=0):
    return datetime(y, mo, d, h, mi, s, tzinfo=LAGOS)


def _fact(dt: datetime) -> str:
    return dates.to_source_fact(dt)


def _sms_utc(dt: datetime) -> str:
    return dates.to_utc_iso(dt)


def _phone(uid: str) -> str:
    return f"080{int(uid):0>8}"


def _setup(_init_db, name="test"):
    win = svc.create_window(name=name, start_time=_dt(2026, 9, 14))
    svc.set_control_override(win["id"], 50.0)
    svc.set_eligible_count(win["id"], 1, {"unsegmented": 1})
    return win


def _add_members(win_id, run_id, user_ids, segment="unsegmented") -> None:
    members = [
        {
            "user_id": uid,
            "segment_id": segment,
            "phone": _phone(uid),
            "run_id": run_id,
            "eligibility_state": {"first_name": f"User{uid}"},
        }
        for uid in user_ids
    ]
    result = svc.add_eligible_users(win_id, members)
    assert result["invalid_phone"] == 0
    return result


def _play(uid, at: datetime, game="Spin", amount=100, file="Sales_x.csv"):
    return {
        "user_id": uid,
        "played_at": _fact(at),
        "game_name": game,
        "amount": float(amount),
        "source_file": file,
        "source_key": f"{uid}|{_fact(at)}|{game}",
    }


def _deposit(uid, at: datetime, file="Deposit_x.csv"):
    return {
        "user_id": uid,
        "deposited_at": _fact(at),
        "source_file": file,
        "source_key": f"{uid}|{_fact(at)}",
    }


def _login(uid, at: datetime, file="Login_x.csv"):
    return {
        "user_id": uid,
        "logged_at": _fact(at),
        "source_file": file,
        "source_key": login_source_key(uid, _fact(at)),
    }


def _welcome(run_id, uid, at: datetime, status="sent"):
    log_sms({
        "user_id": uid, "kind": "welcome", "phone": _phone(uid),
        "status": status, "cycle_id": f"run:{run_id}", "sent_at": _sms_utc(at),
    })


def _export_bytes(window_id: int) -> bytes:
    filename, stream = window_export.build_evaluation_workbook(window_id)
    assert filename.startswith("Afrimillions_")
    assert filename.endswith(f"Window{window_id}_Evaluation.xlsx")
    return stream.getvalue()


def _worksheet(xlsx_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Evaluation" in wb.sheetnames
    return wb["Evaluation"]


def _parse_blocks(xlsx_bytes: bytes) -> list[dict]:
    """Block dicts from the worksheet rows (title, window line, headers, rows)."""
    blocks: list[dict] = []
    current = None
    for row_idx, row in enumerate(_worksheet(xlsx_bytes).iter_rows(values_only=True), start=1):
        first = row[0]
        if isinstance(first, str) and first.startswith("CAMPAIGN EVALUATION:"):
            current = {"title": first, "title_row": row_idx, "campaign": None, "control": None}
            blocks.append(current)
        elif current is None or first is None:
            current = None
            continue
        elif isinstance(first, str) and first.startswith("WINDOW:"):
            current["window_line"] = first
            current["window_row"] = row_idx
        elif first == GROUP_HEADER:
            current["headers"] = list(row)
            current["header_row"] = row_idx
        elif first == "CAMPAIGN":
            current["campaign"] = list(row)
            current["campaign_row"] = row_idx
        elif first == "CONTROL":
            current["control"] = list(row)
            current["control_row"] = row_idx
    return blocks


def _metric(row_values: list, header: str):
    return row_values[1 + METRIC_HEADERS.index(header)]


# ---------------------------------------------------------------------------
# Workbook structure + endpoint
# ---------------------------------------------------------------------------

class TestExportStructure:
    def test_endpoint_returns_downloadable_xlsx(self, client, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)

        resp = client.get(f"/windows/{win['id']}/export")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        disposition = resp.headers["content-disposition"]
        assert "attachment" in disposition
        assert f"Window{win['id']}_Evaluation.xlsx" in disposition
        _worksheet(resp.content)

    def test_endpoint_missing_window_returns_404(self, client, _init_db):
        assert client.get("/windows/999999/export").status_code == 404

    def test_single_block_layout(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)

        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert len(blocks) == 1
        block = blocks[0]
        assert block["title"] == "CAMPAIGN EVALUATION: Unsegmented"
        assert block["window_line"].startswith("WINDOW: test 2026-09-14 00:00:00 to 2026-09-19 23:59:59")
        assert block["window_line"].endswith("| CONTROL: 50%")
        assert block["headers"][0] == GROUP_HEADER
        assert block["headers"][1:] == METRIC_HEADERS
        assert block["campaign"][0] == "CAMPAIGN"
        assert block["control"][0] == "CONTROL"
        # Separator before any next block: next title row is two rows after the
        # control row (control row + one blank separator).
        assert block["campaign_row"] - block["header_row"] == 1
        assert block["control_row"] - block["campaign_row"] == 1

    def test_blank_separator_between_blocks(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, CAMPAIGN_IDS, segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["NewPlayedPlayers"])["run"]["id"]
        _add_members(win["id"], run2, CONTROL_IDS, segment="NewPlayedPlayers")

        blocks = _parse_blocks(_export_bytes(win["id"]))
        # second block's title sits directly below a blank separator row.
        assert blocks[1]["title_row"] - blocks[0]["control_row"] == 2


# ---------------------------------------------------------------------------
# Segment separation
# ---------------------------------------------------------------------------

class TestSegmentBlocks:
    def test_one_segment_produces_one_block(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)
        data = window_report.segment_evaluations(win["id"])
        assert [b["segment_id"] for b in data["segments"]] == ["unsegmented"]
        assert data["segments"][0]["segment_label"] == "Unsegmented"

    def test_multiple_segments_produce_separate_blocks(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, CAMPAIGN_IDS, segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["NewPlayedPlayers"])["run"]["id"]
        _add_members(win["id"], run2, CONTROL_IDS, segment="NewPlayedPlayers")

        data = window_report.segment_evaluations(win["id"])
        assert [b["segment_id"] for b in data["segments"]] == ["unsegmented", "NewPlayedPlayers"]
        assert data["segments"][1]["segment_label"] == (
            "New Played Players (first deposit <= 14 days, has played)"
        )

        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert [b["title"] for b in blocks] == [
            "CAMPAIGN EVALUATION: Unsegmented",
            "CAMPAIGN EVALUATION: New Played Players (first deposit <= 14 days, has played)",
        ]
        unseg, newplayed = [b["campaign"] for b in blocks]
        assert _metric(unseg, "Total_Targeted_Audience") == 3
        assert _metric(newplayed, "Total_Targeted_Audience") == 0

    def test_users_not_merged_across_segments(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, ("1", "2"), segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["StandardLapsing"])["run"]["id"]
        _add_members(win["id"], run2, ("3",), segment="StandardLapsing")
        insert_play_records([
            _play("2", _dt(2026, 9, 14, 10), amount=100),  # unsegmented only
            _play("3", _dt(2026, 9, 14, 11), amount=200),  # StandardLapsing only
        ])

        data = window_report.segment_evaluations(win["id"])
        unseg, lapsing = data["segments"]
        assert unseg["segment_id"] == "unsegmented"
        assert lapsing["segment_id"] == "StandardLapsing"
        assert unseg["audience"] == {"campaign": 2, "control": 0}
        assert lapsing["audience"] == {"campaign": 1, "control": 0}
        # Sales never leak between blocks.
        assert unseg["campaign"]["total_sales"] == 100.0
        assert lapsing["campaign"]["total_sales"] == 200.0

    def test_repeated_same_segment_across_runs_is_one_block_and_not_double_counted(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, ("1",), segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run2, ("2",), segment="unsegmented")
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])

        data = window_report.segment_evaluations(win["id"])
        assert len(data["segments"]) == 1
        block = data["segments"][0]
        assert block["segment_id"] == "unsegmented"
        # user 1 counted once despite the repeated selection; user 2 never played.
        assert block["campaign"]["total_targeted_audience"] == 2
        assert block["campaign"]["total_played_users"] == 1
        assert block["campaign"]["total_sales"] == 100.0

    def test_multiple_runs_different_segments_stay_separate(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, ("1", "2"), segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["NeverDepositedUnder400Hrs"])["run"]["id"]
        _add_members(win["id"], run2, ("51",), segment="NeverDepositedUnder400Hrs")

        data = window_report.segment_evaluations(win["id"])
        blocks = {b["segment_id"]: b for b in data["segments"]}
        assert set(blocks) == {"unsegmented", "NeverDepositedUnder400Hrs"}
        assert blocks["unsegmented"]["audience"] == {"campaign": 2, "control": 0}
        assert blocks["NeverDepositedUnder400Hrs"]["audience"] == {"campaign": 0, "control": 1}


# ---------------------------------------------------------------------------
# Campaign / Control from persisted assignment (read-only export)
# ---------------------------------------------------------------------------

class TestCampaignControl:
    def test_values_reflect_persisted_assignment(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)
        insert_play_records([
            _play("1", _dt(2026, 9, 14, 10), amount=100),
            _play("51", _dt(2026, 9, 14, 10), amount=60),
        ])

        blocks = _parse_blocks(_export_bytes(win["id"]))
        block = blocks[0]
        assert _metric(block["campaign"], "Total_Played_Users") == 1
        assert _metric(block["campaign"], "Total_Sales") == 100.0
        assert _metric(block["control"], "Total_Played_Users") == 1
        assert _metric(block["control"], "Total_Sales") == 60.0

    def test_export_does_not_mutate_state(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])
        insert_deposit_records([_deposit("1", _dt(2026, 9, 14, 9))])
        insert_login_records([_login("1", _dt(2026, 9, 14, 9))])

        before_audience = svc.get_audience(win["id"])
        before_window = svc.get_window(win["id"])
        _export_bytes(win["id"])

        assert svc.get_audience(win["id"]) == before_audience
        assert svc.get_window(win["id"]) == before_window
        assert svc.get_window(win["id"])["status"] == "active"


# ---------------------------------------------------------------------------
# Metric values
# ---------------------------------------------------------------------------

class TestMetrics:
    def _seed_full_scenario(self, win_id, run_id):
        insert_play_records([
            _play("1", _dt(2026, 9, 14, 10, 0), amount=100),  # day 1
            _play("1", _dt(2026, 9, 14, 12, 0), amount=50),   # day 1
            _play("1", _dt(2026, 9, 15, 9, 0), amount=200),   # day 2 -> multi-day
            _play("2", _dt(2026, 9, 14, 9, 0), amount=75),    # day 1, single play
            _play("51", _dt(2026, 9, 14, 9, 0), amount=60),   # control
        ])
        insert_deposit_records([
            _deposit("1", _dt(2026, 9, 14, 9)),   # deposits before her play
            _deposit("3", _dt(2026, 9, 15, 9)),   # deposits, never plays
        ])
        insert_login_records([
            _login("1", _dt(2026, 9, 14, 9)),
            _login("2", _dt(2026, 9, 14, 10)),
            _login("3", _dt(2026, 9, 15, 9)),
        ])
        return run_id

    def test_representative_values(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)
        self._seed_full_scenario(win["id"], run_id)

        data = window_report.segment_evaluations(win["id"])
        camp = data["segments"][0]["campaign"]
        ctrl = data["segments"][0]["control"]

        assert camp["total_targeted_audience"] == 3
        assert camp["total_logged_in_users"] == 3
        assert camp["total_played_users"] == 2
        assert camp["total_deposited_users"] == 2
        assert camp["total_sales"] == 425.0
        assert camp["total_plays"] == 4
        assert camp["login_rate"] == 1.0
        assert camp["play_rate"] == round(2 / 3, 4)
        assert camp["deposit_rate"] == round(2 / 3, 4)
        assert camp["arpu"] == round(425.0 / 3, 4)
        assert camp["arppu"] == 212.5
        assert camp["plays_per_player"] == 2.0
        assert camp["active_days"] == 2
        assert camp["active_days_per_player"] == 1.0
        assert camp["multi_day_players"] == 1
        assert camp["multi_day_player_rate"] == 0.5
        assert camp["deposit_to_play_rate"] == 0.5
        assert camp["login_to_play_rate"] == round(2 / 3, 4)
        assert camp["deposited_no_play_users"] == 1
        assert camp["logged_in_no_play_users"] == 1

        assert ctrl["total_targeted_audience"] == 2
        assert ctrl["total_played_users"] == 1
        assert ctrl["total_sales"] == 60.0
        assert ctrl["total_deposited_users"] == 0
        assert ctrl["total_logged_in_users"] == 0
        assert ctrl["deposit_to_play_rate"] is None
        assert ctrl["login_to_play_rate"] is None

    def test_workbook_values_match_evaluation_data(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS + CONTROL_IDS)
        self._seed_full_scenario(win["id"], run_id)

        data = window_report.segment_evaluations(win["id"])
        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert len(blocks) == len(data["segments"])
        for block_data, sheet_block in zip(data["segments"], blocks):
            for group, row_values in (("campaign", sheet_block["campaign"]),
                                      ("control", sheet_block["control"])):
                for header, key in zip(
                    METRIC_HEADERS, [k for _h, k, _f in window_export.METRIC_COLUMNS]
                ):
                    expected = block_data[group][key]
                    actual = _metric(row_values, header)
                    if expected is None:
                        assert actual == UNAVAILABLE, (group, header)
                    else:
                        assert actual == float(expected), (group, header)

    def test_percent_cells_use_percentage_format(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])

        ws = _worksheet(_export_bytes(win["id"]))
        camp_row = 4
        login_rate_col = 1 + METRIC_HEADERS.index("Login_Rate") + 1  # +Group offset
        play_rate_col = 1 + METRIC_HEADERS.index("Play_Rate") + 1
        arpu_col = 1 + METRIC_HEADERS.index("ARPU") + 1
        assert ws.cell(row=camp_row, column=login_rate_col).number_format == "0.00%"
        assert ws.cell(row=camp_row, column=play_rate_col).number_format == "0.00%"
        assert ws.cell(row=camp_row, column=arpu_col).number_format == "0.00"

    def test_count_cells_hold_integers(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)
        insert_play_records([_play("1", _dt(2026, 9, 14, 10))])

        ws = _worksheet(_export_bytes(win["id"]))
        targeted_col = 1 + METRIC_HEADERS.index("Total_Targeted_Audience") + 1
        played_col = 1 + METRIC_HEADERS.index("Total_Played_Users") + 1
        assert ws.cell(row=4, column=targeted_col).value == 3
        assert ws.cell(row=4, column=played_col).value == 1


# ---------------------------------------------------------------------------
# Edge cases
# ---------------------------------------------------------------------------

class TestEdgeCases:
    def test_group_with_no_members_emits_em_dash_for_rates(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        # Only Control members -> the Campaign group has zero targeted users.
        _add_members(win["id"], run_id, CONTROL_IDS)

        data = window_report.segment_evaluations(win["id"])
        camp = data["segments"][0]["campaign"]
        assert camp["total_targeted_audience"] == 0
        assert camp["login_rate"] is None  # 0/0 stays None, never 0.0
        assert camp["play_rate"] is None
        assert camp["deposit_rate"] is None
        assert camp["arppu"] is None
        assert camp["plays_per_player"] is None

        blocks = _parse_blocks(_export_bytes(win["id"]))
        campaign_row = blocks[0]["campaign"]
        assert _metric(campaign_row, "Total_Targeted_Audience") == 0
        for header in ("Login_Rate", "Play_Rate", "Deposit_Rate", "ARPU",
                       "ARPPU", "Plays_Per_Player", "Active_Days_Per_Player"):
            assert _metric(campaign_row, header) == UNAVAILABLE, header
        # Control keeps real values and rates with a non-zero denominator.
        control_row = blocks[0]["control"]
        assert _metric(control_row, "Total_Targeted_Audience") == 2
        assert _metric(control_row, "Login_Rate") == 0.0

    def test_zero_activity_rates_stay_zero_only_when_denominator_exists(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)  # 3 members, no activity

        data = window_report.segment_evaluations(win["id"])
        camp = data["segments"][0]["campaign"]
        # audience denominator exists -> honest 0.0 for direct rates.
        assert camp["login_rate"] == 0.0
        assert camp["play_rate"] == 0.0
        assert camp["deposit_rate"] == 0.0
        # per-player / interplay denominators are zero -> None (never 0).
        assert camp["arppu"] is None
        assert camp["plays_per_player"] is None
        assert camp["active_days_per_player"] is None
        assert camp["multi_day_player_rate"] is None
        assert camp["deposit_to_play_rate"] is None
        assert camp["login_to_play_rate"] is None

        blocks = _parse_blocks(_export_bytes(win["id"]))
        campaign_row = blocks[0]["campaign"]
        for header in ("ARPPU", "Plays_Per_Player", "Active_Days_Per_Player",
                       "MultiDay_Player_Rate", "Deposit_To_Play_Rate",
                       "Login_To_Play_Rate"):
            assert _metric(campaign_row, header) == UNAVAILABLE
        assert _metric(campaign_row, "Play_Rate") == 0.0

    def test_multiple_runs_full_block(self, _init_db):
        win = _setup(_init_db)
        run1 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run1, CAMPAIGN_IDS, segment="unsegmented")
        run2 = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run2, CONTROL_IDS, segment="unsegmented")

        data = window_report.segment_evaluations(win["id"])
        assert len(data["segments"]) == 1
        assert data["segments"][0]["audience"] == {"campaign": 3, "control": 2}

    def test_empty_window_exports_structural_block_not_a_blank_sheet(self, _init_db):
        win = svc.create_window(name="empty", start_time=_dt(2026, 9, 14))
        svc.set_control_override(win["id"], 50.0)
        svc.set_eligible_count(win["id"], 1, {"unsegmented": 1})

        data = window_report.segment_evaluations(win["id"])
        # No phone-valid audience yet -> the structural Unsegmented block keeps
        # the workbook from being blank; memberships are empty.
        assert len(data["segments"]) == 1
        block = data["segments"][0]
        assert block["segment_id"] == "unsegmented"
        assert block["audience"] == {"campaign": 0, "control": 0}
        assert block["campaign"]["total_targeted_audience"] == 0

        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert len(blocks) == 1
        assert blocks[0]["campaign"][0] == "CAMPAIGN"
        assert blocks[0]["control"][0] == "CONTROL"
        # Counts are honest zeros, per-player/interplay rates stay unavailable.
        assert _metric(blocks[0]["campaign"], "Total_Targeted_Audience") == 0
        assert _metric(blocks[0]["campaign"], "ARPU") == UNAVAILABLE

    def test_active_window_export_works(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])
        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert _metric(blocks[0]["campaign"], "Total_Sales") == 100.0

    def test_ended_window_export_works(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])
        svc.end_window(win["id"])
        assert svc.get_window(win["id"])["status"] == "ended"
        blocks = _parse_blocks(_export_bytes(win["id"]))
        assert _metric(blocks[0]["campaign"], "Total_Sales") == 100.0

    def test_finalized_window_export_is_frozen_and_read_only(self, _init_db):
        win = _setup(_init_db)
        run_id = svc.start_run(win["id"], segments=["unsegmented"])["run"]["id"]
        _add_members(win["id"], run_id, CAMPAIGN_IDS)
        _welcome(run_id, "1", _dt(2026, 9, 14, 9))
        insert_play_records([_play("1", _dt(2026, 9, 14, 10), amount=100)])
        insert_deposit_records([_deposit("1", _dt(2026, 9, 14, 9))])
        insert_login_records([_login("1", _dt(2026, 9, 14, 9))])

        svc.end_window(win["id"])
        finalized = svc.finalize_window(win["id"], now=FINALIZE_AT)
        assert finalized["status"] == "finalized"

        frozen_before = window_report.get_report(win["id"])
        blocks = _parse_blocks(_export_bytes(win["id"]))
        camp = blocks[0]["campaign"]
        assert _metric(camp, "Total_Targeted_Audience") == 3
        assert _metric(camp, "Total_Sales") == 100.0
        assert _metric(camp, "Login_Rate") == round(1 / 3, 4)

        # Repeating the export is deterministic and never mutates the frozen
        # window-level report.
        blocks2 = _parse_blocks(_export_bytes(win["id"]))
        assert blocks2 == blocks
        assert window_report.get_report(win["id"]) == frozen_before
        assert svc.get_window(win["id"])["status"] == "finalized"