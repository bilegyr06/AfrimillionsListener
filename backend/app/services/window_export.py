"""Campaign Window Excel evaluation export (v2 segment-by-segment workbook).

Builds a single `.xlsx` worksheet with one evaluation block per segment held by
the Window's phone-valid audience (see
`app.services.window_report.segment_evaluations`). Each block matches the
recognizable historical layout:

    CAMPAIGN EVALUATION: <SEGMENT LABEL>
    WINDOW: <WINDOW NAME> <EVAL START> to <EVAL END> | CONTROL: <CTRL PCT>%
    Group | Total_Targeted_Audience | ... | LoggedIn_NoPlay_Users
    CAMPAIGN | <metrics>
    CONTROL  | <metrics>
    <blank separator>

Metric values come straight from the established v2 `_group_metrics` engine
with its rate semantics: a rate with a zero denominator is `None` and renders
as an em dash (never zero); a non-zero rate renders as a percentage.
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import dates
from app.services import window_report

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

#: The dimension identifying the two comparison rows (renamed from the
#: historical `ControlGroup` header; see the export contract).
GROUP_HEADER = "Group"

#: (exported column header, _group_metrics key, openpyxl number format).
METRIC_COLUMNS: list[tuple[str, str, str]] = [
    ("Total_Targeted_Audience", "total_targeted_audience", "0"),
    ("Total_LoggedIn_Users", "total_logged_in_users", "0"),
    ("Total_Played_Users", "total_played_users", "0"),
    ("Total_Deposited_Users", "total_deposited_users", "0"),
    ("Total_Sales", "total_sales", "0.00"),
    ("Login_Rate", "login_rate", "0.00%"),
    ("Play_Rate", "play_rate", "0.00%"),
    ("Deposit_Rate", "deposit_rate", "0.00%"),
    ("ARPU", "arpu", "0.00"),
    ("ARPPU", "arppu", "0.00"),
    ("Plays_Per_Player", "plays_per_player", "0.00"),
    ("Active_Days_Per_Player", "active_days_per_player", "0.00"),
    ("MultiDay_Player_Rate", "multi_day_player_rate", "0.00%"),
    ("Deposit_To_Play_Rate", "deposit_to_play_rate", "0.00%"),
    ("Login_To_Play_Rate", "login_to_play_rate", "0.00%"),
    ("Deposited_NoPlay_Users", "deposited_no_play_users", "0"),
    ("LoggedIn_NoPlay_Users", "logged_in_no_play_users", "0"),
]

#: Marker for metrics that are genuinely unavailable (zero denominator / N/A).
UNAVAILABLE = "\u2014"

_TITLE_FONT = Font(bold=True, size=12)
_WINDOW_FONT = Font(italic=True, size=10)
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(fill_type="solid", start_color="D9E1F2", end_color="D9E1F2")
_GROUP_FONT = Font(bold=True)


def _wall_clock(iso_utc: str) -> str:
    """Render a stored UTC instant in Africa/Lagos wall-clock form."""
    lagos = dates.as_business(dates.parse_utc_iso(iso_utc))
    return lagos.strftime("%Y-%m-%d %H:%M:%S")


def _control_text(control_percentage: float | None) -> str:
    if control_percentage is None:
        return UNAVAILABLE
    return f"{control_percentage:g}%"


def _sanitize_name(value: str | None) -> str:
    """A filesystem/header-safe token for the export filename."""
    if not value:
        return "Campaign"
    token = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return token or "Campaign"


def evaluation_filename(window_id: int, window_name: str | None) -> str:
    """Download filename carrying the Window identity and an export indicator."""
    return f"Afrimillions_{_sanitize_name(window_name)}_Window{window_id}_Evaluation.xlsx"


def _write_metric_row(ws, row: int, group: str, metrics: dict, formats: dict) -> None:
    cell = ws.cell(row=row, column=1, value=group)
    cell.font = _GROUP_FONT
    for col, (header, key, number_format) in enumerate(METRIC_COLUMNS, start=2):
        value = metrics.get(key)
        if value is None:
            ws.cell(row=row, column=col, value=UNAVAILABLE)
        elif number_format == "0":
            ws.cell(row=row, column=col, value=int(value))
        else:
            cell = ws.cell(row=row, column=col, value=float(value))
            cell.number_format = formats[key]


def build_evaluation_workbook(
    window_id: int,
) -> tuple[str, io.BytesIO]:
    """Build the segment-by-segment evaluation workbook for a Window.

    Returns (downloadable filename, xlsx bytes stream). Pure read: the
    underlying segment_evaluations calculation never ingests files and never
    mutates Window/Run/audience/report state.
    """
    data = window_report.segment_evaluations(window_id)

    window = data["window"]
    name = window["name"] or f"Window #{window['id']}"
    period_label = (
        f"{_wall_clock(data['evaluation_period']['start'])} to "
        f"{_wall_clock(data['evaluation_period']['end'])}"
    )
    control_text = _control_text(window["control_percentage"])

    format_by_key = {key: fmt for _, key, fmt in METRIC_COLUMNS}

    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"

    width_by_col = {1: max(len(GROUP_HEADER) + 4, 12)}
    for index, (header, _key, _fmt) in enumerate(METRIC_COLUMNS, start=2):
        width_by_col[index] = max(len(header) + 4, 14)

    for index, width in width_by_col.items():
        ws.column_dimensions[get_column_letter(index)].width = min(width, 42)

    row = 1
    for block in data["segments"]:
        ws.cell(row=row, column=1, value=f"CAMPAIGN EVALUATION: {block['segment_label']}").font = _TITLE_FONT
        ws.cell(
            row=row + 1,
            column=1,
            value=(
                f"WINDOW: {name} {period_label} | "
                f"CONTROL: {control_text}"
            ),
        ).font = _WINDOW_FONT

        for col, (header, _key, _fmt) in enumerate(METRIC_COLUMNS, start=2):
            hcell = ws.cell(row=row + 2, column=col, value=header)
            hcell.font = _HEADER_FONT
            hcell.fill = _HEADER_FILL
            hcell.alignment = Alignment(wrap_text=True, vertical="center")
        gcell = ws.cell(row=row + 2, column=1, value=GROUP_HEADER)
        gcell.font = _HEADER_FONT
        gcell.fill = _HEADER_FILL
        gcell.alignment = Alignment(wrap_text=True, vertical="center")

        _write_metric_row(ws, row + 3, "CAMPAIGN", block["campaign"], format_by_key)
        _write_metric_row(ws, row + 4, "CONTROL", block["control"], format_by_key)

        row += 6  # block (5 rows) + blank separator row

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return evaluation_filename(window_id, window["name"]), stream